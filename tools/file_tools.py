"""
文件读取工具 - 让 AI 能够读取本地的 Excel、Word、TXT、CSV 文件
"""
import os
from pathlib import Path
from typing import Optional, Type
from langchain_core.tools import BaseTool
from pydantic import BaseModel, Field

from config import KNOWLEDGE_BASE_DIR, SUPPORTED_FILE_TYPES


# ============================
# 定义工具的输入参数
# ============================

class ReadFileInput(BaseModel):
    """读取文件的参数"""
    file_path: str = Field(description="文件名（如 '销售数据.xlsx'）或相对路径")


class SearchFilesInput(BaseModel):
    """搜索文件的参数"""
    keyword: str = Field(description="要搜索的关键词")
    file_extension: Optional[str] = Field(default=None, description="文件后缀过滤，如 '.xlsx'")


# ============================
# 工具1：列出所有可用文件
# ============================

class ListFilesTool(BaseTool):
    """列出知识库中的所有文件"""
    name: str = "list_knowledge_files"
    description: str = "列出知识库中所有可用的文件，返回文件名、类型和大小。当用户想知道有哪些文件时使用。"

    def _run(self) -> str:
        if not KNOWLEDGE_BASE_DIR.exists():
            return f"知识库文件夹不存在: {KNOWLEDGE_BASE_DIR}"

        files = []
        for f in sorted(KNOWLEDGE_BASE_DIR.rglob("*")):
            if f.is_file():
                rel_path = f.relative_to(KNOWLEDGE_BASE_DIR)
                size = f.stat().st_size
                ftype = SUPPORTED_FILE_TYPES.get(f.suffix.lower(), "其他")
                size_str = f"{size / 1024:.1f}KB" if size < 1024 * 1024 else f"{size / 1024 / 1024:.1f}MB"
                files.append(f"  📄 {rel_path}  [{ftype}]  ({size_str})")

        if not files:
            return f"知识库文件夹为空，请把文件放到: {KNOWLEDGE_BASE_DIR}"

        return f"📁 知识库中共有 {len(files)} 个文件:\n" + "\n".join(files)


# ============================
# 工具2：读取文件内容
# ============================

class ReadFileTool(BaseTool):
    """读取知识库中的文件"""
    name: str = "read_knowledge_file"
    description: str = "读取知识库中指定文件的内容。支持 .txt .csv .xlsx .docx 格式。"
    args_schema: Type[BaseModel] = ReadFileInput

    def _run(self, file_path: str) -> str:
        full_path = KNOWLEDGE_BASE_DIR / file_path

        if not full_path.exists():
            # 尝试模糊匹配
            matches = list(KNOWLEDGE_BASE_DIR.rglob(f"*{file_path}*"))
            if matches:
                full_path = matches[0]
            else:
                return f"❌ 文件不存在: {file_path}\n请先用 list_knowledge_files 查看有哪些文件。"

        suffix = full_path.suffix.lower()

        try:
            if suffix == '.txt':
                return self._read_text(full_path)
            elif suffix == '.csv':
                return self._read_csv(full_path)
            elif suffix in ['.xlsx', '.xls']:
                return self._read_excel(full_path)
            elif suffix == '.docx':
                return self._read_word(full_path)
            else:
                return f"❌ 不支持的文件类型: {suffix}"
        except Exception as e:
            return f"❌ 读取文件出错: {str(e)}"

    def _read_text(self, path: Path) -> str:
        """读取文本文件"""
        import chardet
        with open(path, 'rb') as f:
            raw = f.read()
            encoding = chardet.detect(raw).get('encoding', 'utf-8')

        with open(path, 'r', encoding=encoding) as f:
            content = f.read()

        return f"📄 {path.name}\n{'=' * 40}\n{content}"

    def _read_csv(self, path: Path) -> str:
        """读取 CSV 文件"""
        import pandas as pd
        df = pd.read_csv(path)

        result = f"📊 {path.name}\n"
        result += f"数据量: {df.shape[0]} 行 × {df.shape[1]} 列\n"
        result += f"列名: {', '.join(df.columns.tolist())}\n"
        result += f"{'=' * 40}\n"

        # 显示前30行
        display_rows = min(30, len(df))
        result += df.head(display_rows).to_string(index=False)

        if len(df) > display_rows:
            result += f"\n\n... 还有 {len(df) - display_rows} 行数据"

        return result

    def _read_excel(self, path: Path) -> str:
        """读取 Excel 文件"""
        import pandas as pd

        xls = pd.ExcelFile(path)
        result = f"📗 {path.name}\n"
        result += f"工作表: {', '.join(xls.sheet_names)}\n"

        for sheet in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet)
            display_rows = min(30, len(df))
            result += f"\n--- 工作表: {sheet} ({df.shape[0]}行 × {df.shape[1]}列) ---\n"
            result += f"列名: {', '.join(df.columns.tolist())}\n"
            result += df.head(display_rows).to_string(index=False)

            if len(df) > display_rows:
                result += f"\n... 还有 {len(df) - display_rows} 行"
            result += "\n"

        return result

    def _read_word(self, path: Path) -> str:
        """读取 Word 文档"""
        from docx import Document
        doc = Document(path)

        result = f"📝 {path.name}\n"
        result += f"段落数: {len(doc.paragraphs)}\n"
        result += f"{'=' * 40}\n\n"

        for para in doc.paragraphs:
            if para.text.strip():
                style = para.style.name if para.style else ""
                if 'Heading' in style:
                    result += f"\n## {para.text}\n"
                else:
                    result += para.text + "\n"

        # 读取表格
        if doc.tables:
            result += f"\n--- 表格 ({len(doc.tables)}个) ---\n"
            for i, table in enumerate(doc.tables):
                result += f"\n表格 {i + 1}:\n"
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    result += " | ".join(cells) + "\n"

        return result


# ============================
# 工具3：搜索文件
# ============================

class SearchFilesTool(BaseTool):
    """在知识库中搜索文件"""
    name: str = "search_knowledge_files"
    description: str = "在知识库中按关键词搜索文件名或文件内容。"
    args_schema: Type[BaseModel] = SearchFilesInput

    def _run(self, keyword: str, file_extension: Optional[str] = None) -> str:
        matches = []
        kw = keyword.lower()

        for f in KNOWLEDGE_BASE_DIR.rglob("*"):
            if not f.is_file():
                continue
            if file_extension and f.suffix.lower() != file_extension.lower():
                continue

            # 检查文件名
            name_match = kw in f.name.lower()

            # 检查文件内容（只查文本类文件）
            content_match = False
            if f.suffix.lower() in ['.txt', '.csv']:
                try:
                    with open(f, 'r', encoding='utf-8', errors='ignore') as fh:
                        content = fh.read(50000)
                        content_match = kw in content.lower()
                except:
                    pass

            if name_match or content_match:
                rel = f.relative_to(KNOWLEDGE_BASE_DIR)
                reason = []
                if name_match:
                    reason.append("文件名匹配")
                if content_match:
                    reason.append("内容匹配")
                matches.append(f"  📄 {rel}  ({', '.join(reason)})")

        if not matches:
            return f"没有找到包含 '{keyword}' 的文件"

        return f"找到 {len(matches)} 个匹配文件:\n" + "\n".join(matches[:20])


# ============================
# 获取所有文件工具
# ============================

def get_file_tools():
    """返回所有文件相关工具的列表"""
    return [ListFilesTool(), ReadFileTool(), SearchFilesTool()]
